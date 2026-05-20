import pdfplumber
import re
import os
from PIL import Image
import pytesseract
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'


def parse_invoice_table(path):
    #Парсинг таблиц в счете
    items = []

    try:
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                tables = page.extract_tables()
                for table in tables:
                    for row in table:
                        if not row or not any(row):
                            continue
                        cells = [str(c).replace('\n', ' ').strip() if c else '' for c in row]
                        first_cell = cells[0].lower() if cells[0] else ''
                        skip = ['№', 'no', '', 'итого', 'итог', 'доставка', 'сумма', 'всего', 'наименование']
                        if any(w == first_cell for w in skip):
                            continue
                        name = ''
                        for c in cells:
                            if len(c) > len(name) and not re.match(r'^[\d\s.,%№]+$', c):
                                name = c
                        price = ''
                        for c in reversed(cells):
                            m = re.search(r'([\d\s]+[.,]\d{2})', c)
                            if m:
                                price = m.group(1).replace(' ', '').replace(',', '.')
                                break
                        if name and price:
                            try:
                                items.append({'name': name, 'price': float(price)})
                            except ValueError:
                                print(f"Не удалось преобразовать цену: {price}")
    except Exception as e:
        print(f"Ошибка при парсинге таблиц {path}: {e}")

    return items


def parse_invoice_text(path):
    #Парсинг текстовых счетов
    items = []

    try:
        with pdfplumber.open(path) as pdf:
            full_text = ""
            for page in pdf.pages:
                t = page.extract_text()
                if t:
                    full_text += t + "\n"

        for line in full_text.split('\n'):
            line = line.strip()
            if not line:
                continue
            m = re.search(
                r'(С?\d{3}-\d\s+.+?\d+\s*мм)\s*[-–—]\s*\d+\s*шт.*?=\s*(?:сумма\s*)?([\d\s]+)\s*р',
                line, re.IGNORECASE
            )
            if m:
                try:
                    name = m.group(1).strip()
                    price = float(m.group(2).replace(' ', ''))
                    items.append({'name': name, 'price': price})
                except ValueError:
                    print(f"Не удалось преобразовать цену в строке: {line}")
    except Exception as e:
        print(f"Ошибка при текстовом парсинге {path}: {e}")

    return items


def parse_positions_from_text(text):
   #Парсинг позиций с количеством
    positions = []

    try:
        for line in text.split('\n'):
            line = line.strip()
            if not line:
                continue
            m = re.search(r'(.+?)\s*[-–—]\s*(\d+)\s*(шт|упаковк[аи]|уп|пог\.?\s*метр)', line, re.IGNORECASE)
            if not m:
                continue
            name = m.group(1).strip()
            qty = int(m.group(2))
            unit = m.group(3)


            if 'упаковк' in line or 'уп' in line:
                m2 = re.search(r'Кол-во\s+в\s+упаковке[,:\s]*\s*(\d+)', line, re.IGNORECASE)
                if m2:
                    try:
                        qty = qty * int(m2.group(1))
                        unit = 'шт'
                    except ValueError:
                        print(f"Ошибка при пересчете упаковок для {name}")

            positions.append({'name': name, 'qty': qty, 'unit': unit})
    except Exception as e:
        print(f"Ошибка при парсинге позиций из текста: {e}")

    return positions


def parse_sz(path):
    #Парсинг сметной спецификации
    try:
        text = ""
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                t = page.extract_text()
                if t:
                    text += t + "\n"

        if not text.strip():
            print("Внимание: не удалось извлечь текст из СЗ")
            return []

        return parse_positions_from_text(text)
    except Exception as e:
        print(f"Ошибка при парсинге СЗ {path}: {e}")
        return []


def parse_correction(path):
    #Парсинг файла корректировки
    text = ""

    try:
        if path.lower().endswith('.pdf'):
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    t = page.extract_text()
                    if t:
                        text += t + "\n"
            if not text.strip():
                print("Текст не найден, пробуем OCR...")
                with pdfplumber.open(path) as pdf:
                    for page in pdf.pages:
                        img = page.to_image(resolution=300)
                        text += pytesseract.image_to_string(img.original, lang='rus+eng') + "\n"
        else:
            img = Image.open(path)
            text = pytesseract.image_to_string(img, lang='rus+eng')

        return parse_positions_from_text(text)
    except Exception as e:
        print(f"Ошибка при парсинге корректировки {path}: {e}")
        return []


def apply_correction(positions, corrections):
    """Применение корректировок"""
    try:
        for corr in corrections:
            corr_type = get_type(corr['name'])
            c_nums = get_comparable_nums(corr['name'])
            if not c_nums:
                continue

            for poz in positions:
                poz_type = get_type(poz['name'])
                p_nums = get_comparable_nums(poz['name'])
                if not p_nums:
                    continue

                if corr_type == poz_type and c_nums[0] == p_nums[0]:
                    poz['qty'] = corr['qty']
                    poz['unit'] = corr['unit']

                    if len(c_nums) >= 2 and len(p_nums) >= 2:
                        for i in range(1, min(len(c_nums), len(p_nums))):
                            if c_nums[i] != p_nums[i]:
                                poz['name'] = poz['name'].replace(p_nums[i], c_nums[i])
                    break
    except Exception as e:
        print(f"Ошибка при применении корректировок: {e}")

    return positions


def get_type(name):
    #Определение типа материала
    try:
        n = name.lower()
        if 'плита' in n or 'плиты' in n or 'baswool' in n or 'тизол' in n or 'euro' in n:
            return 'плита'
        if 'круг' in n:
            return 'круг'
        if 'л-' in n or 'лист' in n or 'n-' in n:
            return 'лист'
        if 'уголок' in n:
            return 'уголок'
        if 'швеллер' in n:
            return 'швеллер'
        if 'полоса' in n:
            return 'полоса'
        if 'шнур' in n:
            return 'шнур'
        return ''
    except:
        return ''


def get_nums(name):
    #Извлечение всех чисел из названия
    try:
        return re.findall(r'\d+', name)
    except:
        return []


def get_comparable_nums(name):
    try:
        nums = get_nums(name)
        if not nums:
            return []

        n = name.lower()

        # Для плит берем до 4 чисел
        if 'плита' in n or 'плиты' in n or 'baswool' in n or 'тизол' in n or 'euro' in n:
            return nums[:4] if len(nums) >= 4 else nums

        # Специальные случаи
        if len(nums) >= 3 and nums[0] in ('245', '345'):
            return nums[2:]
        if len(nums) >= 3 and nums[0] == '09':
            return nums[2:]
        if len(nums) >= 3 and nums[0] == '3':
            return nums[1:]

        return nums
    except:
        return []


def get_mark(name):
    try:
        n = name.lower()
        if 'ст3' in n or 'ст 3' in n:
            return 'ст3'
        if '09г2с-15' in n:
            return '09Г2С-15'
        if '09г2с' in n:
            return '09Г2С'
        if 'с355' in n:
            return 'С355'
        if 'baswool' in n:
            return 'BASWOOL'
        if 'тизол' in n:
            return 'ТИЗОЛ'
        if 'с245' in n:
            return 'ст3'
        if 'с345' in n:
            return '09Г2С'
        return ''
    except:
        return ''


def find_match(poz, items):
    try:
        poz_type = get_type(poz['name'])
        poz_nums = get_comparable_nums(poz['name'])
        best = None
        best_score = -1000

        for item in items:
            item_type = get_type(item['name'])
            item_nums = get_comparable_nums(item['name'])

            if poz_type == 'плита':
                if item_type != 'плита':
                    continue
                if not item_nums or not poz_nums:
                    continue
                score = 0
                if len(poz_nums) >= 2 and len(item_nums) >= 2:
                    if poz_nums[-2:] == item_nums[-2:]:
                        score += 30
                    elif poz_nums[-1] == item_nums[-1]:
                        score += 15
                    else:
                        score -= 50
                if score > best_score:
                    best_score = score
                    best = item
                continue
            if poz_type == 'шнур':
                if 'шнур' in item_type:
                    best = item
                    break
                continue
            if poz_type and item_type and poz_type != item_type:
                continue

            score = 0
            if poz_nums and item_nums:
                if poz_nums[0] == item_nums[0]:
                    score += 10
                else:
                    score -= 100
                if len(poz_nums) >= 2 and len(item_nums) >= 2:
                    if poz_nums[1] == item_nums[1]:
                        score += 5
                    else:
                        score -= 20
                if len(poz_nums) >= 3 and len(item_nums) >= 3:
                    if poz_nums[2] == item_nums[2]:
                        score += 3

            if score > best_score:
                best_score = score
                best = item

        if best:
            if poz_type == 'плита' and best_score >= 10:
                return best
            if poz_type != 'плита' and best_score >= 5:
                return best
        return None
    except Exception as e:
        print(f"Ошибка при поиске соответствия: {e}")
        return None


def save_excel(positions, supplier_names, output_dir):
    try:
        # Создаем папку если её нет
        os.makedirs(output_dir, exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = "Таблица аналогов"

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3 + len(supplier_names) * 2)
        ws.cell(row=1, column=1, value="Список МСЦ 98-26").font = Font(bold=True, size=12)
        ws.cell(row=1, column=1).alignment = Alignment(horizontal='center')

        headers = ['№', 'Наименование', 'Кол-во']
        for s in supplier_names:
            headers.append(f'{s}\nСумма')
            headers.append(f'{s}\nИзм.')

        for col, h in enumerate(headers, 1):
            c = ws.cell(row=3, column=col, value=h)
            c.font = Font(bold=True)
            c.fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        thin = Border(left=Side(style='thin'), right=Side(style='thin'),
                      top=Side(style='thin'), bottom=Side(style='thin'))

        for i, p in enumerate(positions):
            r = i + 4
            ws.cell(row=r, column=1, value=i + 1)
            ws.cell(row=r, column=2, value=p['name'])
            ws.cell(row=r, column=3, value=f"{p['qty']} {p['unit']}")

            for j, s in enumerate(supplier_names):
                col_sum = 4 + j * 2
                col_chg = 5 + j * 2
                price = p.get(s + '_price', None)
                analog = p.get(s + '_analog', '')
                ws.cell(row=r, column=col_sum, value=price if price else '—')
                ws.cell(row=r, column=col_chg, value=analog if analog else '—')

        last_row = len(positions) + 3
        tr = last_row + 2
        ws.cell(row=tr, column=2, value='ИТОГО').font = Font(bold=True)
        for j in range(len(supplier_names)):
            cl = get_column_letter(4 + j * 2)
            ws.cell(row=tr, column=4 + j * 2, value=f'=SUM({cl}4:{cl}{last_row})').font = Font(bold=True)

        for row in ws.iter_rows(min_row=3, max_row=tr, max_col=3 + len(supplier_names) * 2):
            for cell in row:
                cell.border = thin
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                if cell.column == 2:  # Наименование выравниваем влево
                    cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

        ws.column_dimensions['A'].width = 5
        ws.column_dimensions['B'].width = 45
        ws.column_dimensions['C'].width = 12
        for i in range(len(supplier_names)):
            ws.column_dimensions[get_column_letter(4 + i * 2)].width = 16
            ws.column_dimensions[get_column_letter(5 + i * 2)].width = 22

        path1 = os.path.join(output_dir, 'Таблица_аналогов.xlsx')
        wb.save(path1)

        wb2 = Workbook()
        ws2 = wb2.active
        ws2.title = "Согласование замен"

        ws2.merge_cells('A1:G1')
        ws2.cell(row=1, column=1, value="РЕШЕНИЕ о согласовании замены №240").font = Font(bold=True, size=12)
        ws2.cell(row=1, column=1).alignment = Alignment(horizontal='center')
        ws2.merge_cells('A2:G2')
        ws2.cell(row=2, column=1, value="от 17.03.2026").alignment = Alignment(horizontal='center')

        headers2 = ['№', 'Заказ поставщика', 'Заказ/Система', 'в КД', 'Кол-во к закупке', 'Аналог', 'Комментарий']
        for col, h in enumerate(headers2, 1):
            c = ws2.cell(row=4, column=col, value=h)
            c.font = Font(bold=True)
            c.fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')

        row_num = 5
        counter = 1
        for p in positions:
            analogs = []
            for s in supplier_names:
                a = p.get(s + '_analog', '')
                if a:
                    analogs.append(f"{s}: {a}")

            if analogs:
                ws2.cell(row=row_num, column=1, value=counter)
                ws2.cell(row=row_num, column=2, value='UPP00001623')
                ws2.cell(row=row_num, column=3, value='2575 МСЦ 98 система выхлопа с шумоглушения')
                ws2.cell(row=row_num, column=4, value=p['name'])
                ws2.cell(row=row_num, column=5, value=f"{p['qty']} {p['unit']}")
                ws2.cell(row=row_num, column=6, value='; '.join(analogs))
                ws2.cell(row=row_num, column=7, value='V')
                row_num += 1
                counter += 1

        for row in ws2.iter_rows(min_row=4, max_row=row_num - 1, max_col=7):
            for cell in row:
                cell.border = thin
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                if cell.column in [4, 6]:  # Наименование и аналог выравниваем влево
                    cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

        ws2.column_dimensions['A'].width = 5
        ws2.column_dimensions['B'].width = 16
        ws2.column_dimensions['C'].width = 28
        ws2.column_dimensions['D'].width = 42
        ws2.column_dimensions['E'].width = 16
        ws2.column_dimensions['F'].width = 40
        ws2.column_dimensions['G'].width = 13

        path2 = os.path.join(output_dir, 'Лист_согласования_замен.xlsx')
        wb2.save(path2)

        return path1, path2

    except Exception as e:
        print(f"Ошибка при сохранении Excel: {e}")
        return None, None

path_sz = input("\nПуть к СЗ (.pdf): ").strip().strip('"')
positions = parse_sz(path_sz)
print(f"СЗ: {len(positions)} позиций")

path_corr = input("\nПуть к корректировке (.pdf/.png) или Enter: ").strip().strip('"')
if path_corr and os.path.isfile(path_corr):
    corrections = parse_correction(path_corr)
    print(f"Корректировка: {len(corrections)} позиций")
    positions = apply_correction(positions, corrections)
    print("Корректировка применена")
else:
    print("Без корректировки")

suppliers = {}
print("\n" + "-" * 40)
print("Введите поставщиков (пустое имя - завершить)")
print("-" * 40)

while True:
    name = input("\nПоставщик: ").strip()
    if not name:
        if suppliers:
            break
        else:
            print("Нужно добавить хотя бы одного поставщика!")
            continue

    path = input(f"Файл для {name}: ").strip().strip('"')

    items = parse_invoice_table(path)
    if len(items) == 0:
        items = parse_invoice_text(path)

    if items:
        print(f"  {name}: {len(items)} позиций")
        suppliers[name] = items
    else:
        print(f"  {name}: ничего не найдено, пропускаем")

print("\nПоиск соответствий...")
for name, items in suppliers.items():
    print(f"  Обработка {name}...")
    for poz in positions:
        match = find_match(poz, items)
        if match:
            poz[name + '_price'] = match['price']
            poz[name + '_analog'] = get_mark(match['name'])

path1, path2 = save_excel(positions, list(suppliers.keys()), 'output')

print(f"\nГотово!")
print(f"Таблица аналогов: {path1}")
print(f"Лист согласования: {path2}")

try:
    os.startfile('output')
except:
    print(f"Результаты в папке: {os.path.abspath('output')}")